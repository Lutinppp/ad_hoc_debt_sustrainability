#!/usr/bin/env python3
"""Standalone IMF series exporter."""

from __future__ import annotations

import argparse
import os
from dataclasses import dataclass
from typing import Dict, Iterable, List, Tuple

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

COUNTRY_MAP: Dict[str, Tuple[str, str]] = {
    "EZ": ("EURO", "Euro Area"),
    "FR": ("FRA", "France"),
    "DE": ("DEU", "Germany"),
    "IT": ("ITA", "Italy"),
    "ES": ("ESP", "Spain"),
    "US": ("USA", "United States"),
    "UK": ("GBR", "United Kingdom"),
}

INDICATORS = {
    "debt": "GGXWDG_NGDP",
    "primary_balance": "GGXONLB_G01_GDP_PT",
    "gdp_level_primary": "NGDP",
    "gdp_level_fallback": "NGDPD",
}


@dataclass(frozen=True)
class IMFConfig:
    base_url: str = "https://www.imf.org/external/datamapper/api/v1"
    weo_sdmx_base_url: str = "https://api.imf.org/external/sdmx/3.0"
    weo_sdmx_version: str = "9.0.0"
    timeout_seconds: int = 60
    api_key: str = "6d2a81abe2884cc38950d3391e337860"
    bearer_token: str = ""

    @classmethod
    def from_env(cls) -> "IMFConfig":
        return cls(
            base_url=os.getenv("IMF_BASE_URL", cls.base_url),
            weo_sdmx_base_url=os.getenv("IMF_WEO_SDMX_BASE_URL", cls.weo_sdmx_base_url),
            weo_sdmx_version=os.getenv("IMF_WEO_SDMX_VERSION", cls.weo_sdmx_version),
            timeout_seconds=int(os.getenv("IMF_TIMEOUT_SECONDS", str(cls.timeout_seconds))),
            api_key=os.getenv("IMF_API_KEY", cls.api_key),
            bearer_token=os.getenv("IMF_BEARER_TOKEN", cls.bearer_token),
        )

    def headers(self) -> Dict[str, str]:
        headers: Dict[str, str] = {"Accept": "application/json"}
        if self.api_key:
            headers["x-api-key"] = self.api_key
        if self.bearer_token:
            headers["Authorization"] = f"Bearer {self.bearer_token}"
        return headers


@dataclass(frozen=True)
class SeriesRequest:
    indicator: str
    countries: Iterable[str]
    start_year: int
    end_year: int


@dataclass
class SeriesExportInputs:
    start_year: int
    end_year: int
    output_path: str


class IMFClient:
    def __init__(self, config: IMFConfig) -> None:
        self.config = config

    def fetch_indicator(self, indicator: str) -> Dict[str, Dict[int, float]]:
        url = f"{self.config.base_url}/{indicator}"
        response = requests.get(
            url,
            headers=self.config.headers(),
            timeout=self.config.timeout_seconds,
        )
        response.raise_for_status()
        payload = response.json()

        values = payload.get("values", {}).get(indicator)
        if not isinstance(values, dict):
            raise ValueError(f"Indicator '{indicator}' not found in IMF response.")

        cleaned: Dict[str, Dict[int, float]] = {}
        for country_code, year_map in values.items():
            if not isinstance(year_map, dict):
                continue
            parsed_years: Dict[int, float] = {}
            for year_str, value in year_map.items():
                try:
                    year = int(year_str)
                    parsed_years[year] = float(value)
                except (TypeError, ValueError):
                    continue
            if parsed_years:
                cleaned[country_code] = parsed_years
        return cleaned

    def fetch_series(self, req: SeriesRequest) -> Dict[str, Dict[int, float]]:
        all_values = self.fetch_indicator(req.indicator)
        year_filtered: Dict[str, Dict[int, float]] = {}
        wanted = set(req.countries)
        for country_code, year_map in all_values.items():
            if country_code not in wanted:
                continue
            year_filtered[country_code] = {
                year: value
                for year, value in year_map.items()
                if req.start_year <= year <= req.end_year
            }
        return year_filtered

    def fetch_weo_gdp_series(
        self,
        countries: List[str],
        start_year: int,
        end_year: int,
        primary_indicator: str,
        fallback_indicator: str,
    ) -> Tuple[Dict[str, Dict[int, float]], Dict[str, str]]:
        countries_path = "+".join(countries)
        indicators_path = f"{primary_indicator}+{fallback_indicator}"
        url = (
            f"{self.config.weo_sdmx_base_url}/data/dataflow/IMF.RES/WEO/"
            f"{self.config.weo_sdmx_version}/{countries_path}.{indicators_path}.*"
        )
        time_filter = f"ge:{start_year}-01-01+le:{end_year}-12-31"
        query = (
            f"c[TIME_PERIOD]={time_filter}"
            "&attributes=all"
            "&detail=full"
            "&includeHistory=true"
            "&limit=10000"
        )

        response = requests.get(
            f"{url}?{query}",
            headers={"Accept": "application/json"},
            timeout=self.config.timeout_seconds,
        )
        response.raise_for_status()
        payload = response.json()

        parsed = _parse_sdmx_weo_payload(payload)
        by_country: Dict[str, Dict[int, float]] = {}
        indicator_used: Dict[str, str] = {}

        for country in countries:
            primary_series = parsed.get(country, {}).get(primary_indicator, {})
            fallback_series = parsed.get(country, {}).get(fallback_indicator, {})
            if primary_series:
                by_country[country] = primary_series
                indicator_used[country] = primary_indicator
            elif fallback_series:
                by_country[country] = fallback_series
                indicator_used[country] = fallback_indicator
            else:
                by_country[country] = {}
                indicator_used[country] = ""

        return by_country, indicator_used


def _dimension_value_id(dimension: Dict, index: int) -> str | None:
    values = dimension.get("values", [])
    if not isinstance(values, list) or index < 0 or index >= len(values):
        return None
    value = values[index]
    if not isinstance(value, dict):
        return None
    return value.get("id") or value.get("value")


def _parse_sdmx_weo_payload(payload: Dict) -> Dict[str, Dict[str, Dict[int, float]]]:
    parsed: Dict[str, Dict[str, Dict[int, float]]] = {}

    structures = payload.get("data", {}).get("structures", [])
    datasets = payload.get("data", {}).get("dataSets", [])
    if not structures or not datasets:
        return parsed

    dimensions = structures[0].get("dimensions", {})
    series_dims = dimensions.get("series", [])
    obs_dims = dimensions.get("observation", [])
    series_data = datasets[0].get("series", {})
    if not series_dims or not obs_dims or not isinstance(series_data, dict):
        return parsed

    country_index = next((i for i, dim in enumerate(series_dims) if dim.get("id") in ("COUNTRY", "REF_AREA")), None)
    indicator_index = next((i for i, dim in enumerate(series_dims) if dim.get("id") == "INDICATOR"), None)
    if country_index is None or indicator_index is None:
        return parsed

    time_dim = obs_dims[0]

    for series_key, series_values in series_data.items():
        try:
            key_parts = [int(part) for part in series_key.split(":")]
        except ValueError:
            continue

        if country_index >= len(key_parts) or indicator_index >= len(key_parts):
            continue

        country_code = _dimension_value_id(series_dims[country_index], key_parts[country_index])
        indicator_code = _dimension_value_id(series_dims[indicator_index], key_parts[indicator_index])
        if not country_code or not indicator_code:
            continue

        observations = series_values.get("observations", {}) if isinstance(series_values, dict) else {}
        if not isinstance(observations, dict):
            continue

        for obs_key, obs_value in observations.items():
            try:
                time_index = int(obs_key.split(":")[0])
            except ValueError:
                continue

            year_text = _dimension_value_id(time_dim, time_index)
            if not year_text:
                continue

            try:
                year = int(str(year_text)[:4])
            except ValueError:
                continue

            if not isinstance(obs_value, list) or not obs_value:
                continue

            try:
                value = float(obs_value[0])
            except (TypeError, ValueError):
                continue

            parsed.setdefault(country_code, {}).setdefault(indicator_code, {})[year] = value

    return parsed


def _compute_growth(levels: Dict[int, float], years: List[int]) -> Dict[int, float]:
    growth: Dict[int, float] = {}
    for year in years:
        if year not in levels or (year - 1) not in levels:
            continue
        prev = levels[year - 1]
        if prev == 0:
            continue
        growth[year] = ((levels[year] - prev) / prev) * 100.0
    return growth


def _write_country_sheet(
    wb: Workbook,
    sheet_name: str,
    country_label: str,
    years: List[int],
    debt: Dict[int, float],
    primary_balance: Dict[int, float],
    gdp_level: Dict[int, float],
    growth: Dict[int, float],
    gdp_indicator_used: str,
) -> None:
    ws = wb.create_sheet(title=sheet_name)

    headers = [
        "Year",
        "Gross debt (% GDP)",
        "Primary balance (% GDP)",
        f"GDP level ({gdp_indicator_used or 'N/A'})",
        "Nominal GDP growth (%)",
    ]

    ws["A1"] = country_label
    ws["A1"].font = Font(bold=True, size=12)

    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=column_index, value=header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")

    for row_index, year in enumerate(years, start=3):
        ws.cell(row=row_index, column=1, value=year)
        ws.cell(row=row_index, column=2, value=debt.get(year))
        ws.cell(row=row_index, column=3, value=primary_balance.get(year))
        ws.cell(row=row_index, column=4, value=gdp_level.get(year))
        ws.cell(row=row_index, column=5, value=growth.get(year))

    for column in ("A", "B", "C", "D", "E"):
        ws.column_dimensions[column].width = 24

    for row_index in range(3, len(years) + 3):
        ws.cell(row=row_index, column=2).number_format = "0.0"
        ws.cell(row=row_index, column=3).number_format = "0.0"
        ws.cell(row=row_index, column=4).number_format = "0.0"
        ws.cell(row=row_index, column=5).number_format = "0.0"

    ws.freeze_panes = "A3"


def run_series_export(inputs: SeriesExportInputs, config: IMFConfig | None = None) -> str:
    if inputs.end_year < inputs.start_year:
        raise ValueError("end_year must be greater than or equal to start_year")

    cfg = config or IMFConfig.from_env()
    client = IMFClient(cfg)

    fetch_start = inputs.start_year - 1
    visible_years = list(range(inputs.start_year, inputs.end_year + 1))
    fetch_years = list(range(fetch_start, inputs.end_year + 1))
    target_codes = [code for code, _ in COUNTRY_MAP.values()]

    debt_data = client.fetch_series(
        SeriesRequest(
            indicator=INDICATORS["debt"],
            countries=target_codes,
            start_year=inputs.start_year,
            end_year=inputs.end_year,
        )
    )
    pb_data = client.fetch_series(
        SeriesRequest(
            indicator=INDICATORS["primary_balance"],
            countries=target_codes,
            start_year=inputs.start_year,
            end_year=inputs.end_year,
        )
    )

    gdp_data, gdp_indicator_by_country = client.fetch_weo_gdp_series(
        countries=target_codes,
        start_year=fetch_start,
        end_year=inputs.end_year,
        primary_indicator=INDICATORS["gdp_level_primary"],
        fallback_indicator=INDICATORS["gdp_level_fallback"],
    )

    missing_gdp_countries = [code for code in target_codes if not gdp_data.get(code)]
    if missing_gdp_countries:
        try:
            dm_primary = client.fetch_series(
                SeriesRequest(
                    indicator=INDICATORS["gdp_level_primary"],
                    countries=missing_gdp_countries,
                    start_year=fetch_start,
                    end_year=inputs.end_year,
                )
            )
        except ValueError:
            dm_primary = {}
        dm_fallback = client.fetch_series(
            SeriesRequest(
                indicator=INDICATORS["gdp_level_fallback"],
                countries=missing_gdp_countries,
                start_year=fetch_start,
                end_year=inputs.end_year,
            )
        )
        for country in missing_gdp_countries:
            if dm_primary.get(country):
                gdp_data[country] = dm_primary[country]
                gdp_indicator_by_country[country] = f"{INDICATORS['gdp_level_primary']} (DataMapper)"
            elif dm_fallback.get(country):
                gdp_data[country] = dm_fallback[country]
                gdp_indicator_by_country[country] = f"{INDICATORS['gdp_level_fallback']} (DataMapper)"

    wb = Workbook()
    wb.remove(wb.active)

    for short_code, (imf_code, country_label) in COUNTRY_MAP.items():
        debt = debt_data.get(imf_code, {})
        primary_balance = pb_data.get(imf_code, {})
        gdp_level = gdp_data.get(imf_code, {})
        growth = _compute_growth(gdp_level, fetch_years)

        _write_country_sheet(
            wb,
            sheet_name=short_code,
            country_label=country_label,
            years=visible_years,
            debt=debt,
            primary_balance=primary_balance,
            gdp_level=gdp_level,
            growth=growth,
            gdp_indicator_used=gdp_indicator_by_country.get(imf_code, ""),
        )

    info = wb.create_sheet(title="INFO")
    info["A1"] = "Data source"
    info["B1"] = "IMF DataMapper API"
    info["A2"] = "Debt indicator"
    info["B2"] = INDICATORS["debt"]
    info["A3"] = "Primary balance indicator"
    info["B3"] = INDICATORS["primary_balance"]
    info["A4"] = "GDP source"
    info["B4"] = f"IMF SDMX WEO {cfg.weo_sdmx_version}"
    info["A5"] = "GDP indicators"
    info["B5"] = f"{INDICATORS['gdp_level_primary']} with fallback to {INDICATORS['gdp_level_fallback']}"
    info["A6"] = "GDP fallback when SDMX missing"
    info["B6"] = "DataMapper"

    wb.save(inputs.output_path)
    return inputs.output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export IMF debt, primary balance, GDP, and growth series to one Excel sheet per geography."
    )
    parser.add_argument(
        "--start-year",
        type=int,
        default=2020,
        help="First year to export (default: 2020)",
    )
    parser.add_argument(
        "--end-year",
        type=int,
        default=2030,
        help="Last year to export (default: 2030)",
    )
    parser.add_argument(
        "--output",
        default="imf_series_export.xlsx",
        help="Output workbook path (default: imf_series_export.xlsx)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    out_path = run_series_export(
        SeriesExportInputs(start_year=args.start_year, end_year=args.end_year, output_path=args.output)
    )
    print(f"Workbook generated: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())