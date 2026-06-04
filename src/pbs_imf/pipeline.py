"""Debt sustainability pipeline driven by IMF API data."""

from __future__ import annotations

from dataclasses import dataclass
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import IMFConfig
from .imf_client import IMFClient, SeriesRequest

COUNTRY_MAP: Dict[str, Tuple[str, str]] = {
    "EZ": ("EURO", "Euro Area"),
    "FR": ("FRA", "France"),
    "DE": ("DEU", "Germany"),
    "IT": ("ITA", "Italy"),
    "ES": ("ESP", "Spain"),
    "US": ("USA", "United States"),
    "UK": ("GBR", "United Kingdom"),
}

INDICATORS = {
    "debt": "GGXWDG_NGDP",
    "primary_balance": "GGXONLB_G01_GDP_PT",
    # Preferred nominal GDP in national currency is NGDP, but DataMapper commonly serves NGDPD.
    "gdp_level_primary": "NGDP",
    "gdp_level_fallback": "NGDPD",
}

ZERO_FLOW_ROW_LAYOUT = {
    1: "Country",
    2: "Year",
    3: "",
    4: "General government gross debt (% GDP)",
    5: "Projected debt stock (% GDP)",
    6: "Primary net lending/borrowing (% GDP)",
    7: "Nominal GDP growth (%)",
    8: "Other identified flows (% GDP) - set to 0",
    9: "Implied effective rate (%)",
    10: "Debt-stabilizing primary balance differential (pp)",
    11: "Primary balance minus differential (pp)",
}

CALIBRATED_FLOW_ROW_LAYOUT = {
    **ZERO_FLOW_ROW_LAYOUT,
    8: "Computed other identified flows (% GDP)",
}

BLOCK_HEIGHT = max(ZERO_FLOW_ROW_LAYOUT) + 2


@dataclass
class PipelineInputs:
    start_year: int
    end_year: int
    output_path: str


def _compute_growth(levels: Dict[int, float], years: List[int]) -> Dict[int, float]:
    growth: Dict[int, float] = {}
    for year in years:
        if (year not in levels) or ((year - 1) not in levels):
            continue
        prev = levels[year - 1]
        if prev == 0:
            continue
        growth[year] = ((levels[year] - prev) / prev) * 100.0
    return growth


def _build_rows_for_flows(
    years: List[int],
    debt: Dict[int, float],
    primary_balance: Dict[int, float],
    growth: Dict[int, float],
    other_flows: Dict[int, float],
) -> Dict[str, Dict[int, float]]:
    projected: Dict[int, float] = {}
    implied_rate: Dict[int, float] = {}
    differential: Dict[int, float] = {}
    pb_minus_diff: Dict[int, float] = {}

    if years:
        first_year = years[0]
        if first_year in debt:
            projected[first_year] = debt[first_year]

    for year in years[1:]:
        prev_year = year - 1
        prev_proj = projected.get(prev_year)
        g = growth.get(year)
        pb = primary_balance.get(year)
        oif = other_flows.get(year, 0.0)

        d_t = debt.get(year)
        d_prev = debt.get(prev_year)
        if d_t is not None and d_prev not in (None, 0.0) and pb is not None and g is not None:
            implied_rate[year] = ((((d_t + pb - oif) / d_prev) * (1 + g * 0.01)) - 1) * 100.0

        i_rate = implied_rate.get(year)
        if prev_proj is not None and g is not None and pb is not None and i_rate is not None:
            projected[year] = prev_proj * (1 + i_rate / 100.0) / (1 + g / 100.0) - pb + oif

        if d_prev is not None and i_rate is not None and g is not None:
            differential[year] = d_prev * (i_rate - g) * 0.01

        diff = differential.get(year)
        if pb is not None and diff is not None:
            pb_minus_diff[year] = pb - diff

    return {
        "debt": debt,
        "projected": projected,
        "primary_balance": primary_balance,
        "growth": growth,
        "other_flows": other_flows,
        "implied_rate": implied_rate,
        "differential": differential,
        "pb_minus_diff": pb_minus_diff,
    }


def _build_zero_flow_rows(
    years: List[int],
    debt: Dict[int, float],
    primary_balance: Dict[int, float],
    growth: Dict[int, float],
) -> Dict[str, Dict[int, float]]:
    return _build_rows_for_flows(
        years,
        debt=debt,
        primary_balance=primary_balance,
        growth=growth,
        other_flows={y: 0.0 for y in years},
    )


def _build_calibrated_flow_rows(
    years: List[int],
    debt: Dict[int, float],
    primary_balance: Dict[int, float],
    growth: Dict[int, float],
    baseline_rows: Dict[str, Dict[int, float]],
) -> Dict[str, Dict[int, float]]:
    other_flows: Dict[int, float] = {}
    if years:
        other_flows[years[0]] = 0.0

    projected = baseline_rows["projected"]
    for year in years[1:]:
        prev_year = year - 1
        prev_proj = projected.get(prev_year)
        g = growth.get(year)
        pb = primary_balance.get(year)
        d_t = debt.get(year)
        i_rate = baseline_rows["implied_rate"].get(year)
        if prev_proj is None or g is None or pb is None or d_t is None or i_rate is None:
            continue
        flow_adjustment = d_t - (prev_proj * (1 + i_rate / 100.0) / (1 + g / 100.0) - pb)
        other_flows[year] = 0.0 if abs(flow_adjustment) < 1e-9 else flow_adjustment

    return _build_rows_for_flows(
        years,
        debt=debt,
        primary_balance=primary_balance,
        growth=growth,
        other_flows=other_flows,
    )


def _write_block(
    ws,
    start_row: int,
    title: str,
    years: List[int],
    rows: Dict[str, Dict[int, float]],
    row_layout: Dict[int, str],
) -> None:
    ws.cell(row=start_row, column=1, value=title)
    for offset, label in row_layout.items():
        row_number = start_row + offset - 1
        ws.cell(row=row_number, column=1, value=label if offset != 1 else title)

    for idx, year in enumerate(years, start=2):
        ws.cell(row=start_row + 1, column=idx, value=year)
        ws.cell(row=start_row + 3, column=idx, value=rows["debt"].get(year))
        ws.cell(row=start_row + 4, column=idx, value=rows["projected"].get(year))
        ws.cell(row=start_row + 5, column=idx, value=rows["primary_balance"].get(year))
        ws.cell(row=start_row + 6, column=idx, value=rows["growth"].get(year))
        ws.cell(row=start_row + 7, column=idx, value=rows["other_flows"].get(year))
        ws.cell(row=start_row + 8, column=idx, value=rows["implied_rate"].get(year))
        ws.cell(row=start_row + 9, column=idx, value=rows["differential"].get(year))
        ws.cell(row=start_row + 10, column=idx, value=rows["pb_minus_diff"].get(year))


def _write_sheet(
    wb: Workbook,
    sheet_name: str,
    country_label: str,
    years: List[int],
    blocks: List[Tuple[Dict[int, str], Dict[str, Dict[int, float]]]],
) -> None:
    ws = wb.create_sheet(title=sheet_name)

    for block_index, (row_layout, rows) in enumerate(blocks):
        start_row = 1 + (block_index * BLOCK_HEIGHT)
        title = country_label if block_index == 0 else f"{country_label} - calibrated flow scenario"
        _write_block(ws, start_row=start_row, title=title, years=years, rows=rows, row_layout=row_layout)

    # Styling to mimic a report-like worksheet.
    header_fill = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    label_font = Font(bold=True)

    for block_index, _ in enumerate(blocks):
        start_row = 1 + (block_index * BLOCK_HEIGHT)
        ws.cell(row=start_row, column=1).font = Font(bold=True, size=12)
        ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="left")

        for c in range(2, len(years) + 2):
            cell = ws.cell(row=start_row + 1, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for r in range(start_row, start_row + max(ZERO_FLOW_ROW_LAYOUT)):
            ws.cell(row=r, column=1).font = label_font

        for r in range(start_row + 3, start_row + max(ZERO_FLOW_ROW_LAYOUT)):
            for c in range(2, len(years) + 2):
                ws.cell(row=r, column=c).number_format = "0.0"

    ws.column_dimensions["A"].width = 58
    ws.freeze_panes = "B3"

    # Hide lag year column to match the VBA behavior.
    ws.column_dimensions["B"].hidden = True


def run_pipeline(inputs: PipelineInputs, config: IMFConfig | None = None) -> str:
    if inputs.end_year < inputs.start_year:
        raise ValueError("end_year must be greater than or equal to start_year")

    cfg = config or IMFConfig.from_env()
    client = IMFClient(cfg)

    lag_year = inputs.start_year - 1
    fetch_start = lag_year - 1  # needed to compute growth for lag year if available
    years = list(range(lag_year, inputs.end_year + 1))

    target_codes = [code for code, _ in COUNTRY_MAP.values()]

    debt_data = client.fetch_series(
        SeriesRequest(
            indicator=INDICATORS["debt"],
            countries=target_codes,
            start_year=fetch_start,
            end_year=inputs.end_year,
        )
    )
    pb_data = client.fetch_series(
        SeriesRequest(
            indicator=INDICATORS["primary_balance"],
            countries=target_codes,
            start_year=fetch_start,
            end_year=inputs.end_year,
        )
    )

    gdp_indicator_used = INDICATORS["gdp_level_primary"]
    try:
        gdp_data = client.fetch_series(
            SeriesRequest(
                indicator=gdp_indicator_used,
                countries=target_codes,
                start_year=fetch_start,
                end_year=inputs.end_year,
            )
        )
    except ValueError:
        gdp_data = {}

    if not any(gdp_data.get(code) for code in target_codes):
        gdp_indicator_used = INDICATORS["gdp_level_fallback"]
        gdp_data = client.fetch_series(
            SeriesRequest(
                indicator=gdp_indicator_used,
                countries=target_codes,
                start_year=fetch_start,
                end_year=inputs.end_year,
            )
        )

    wb = Workbook()
    wb.remove(wb.active)

    for short_code, (imf_code, country_label) in COUNTRY_MAP.items():
        debt = debt_data.get(imf_code, {})
        pb = pb_data.get(imf_code, {})
        gdp_levels = gdp_data.get(imf_code, {})

        growth = _compute_growth(gdp_levels, years)
        zero_flow_rows = _build_zero_flow_rows(years, debt=debt, primary_balance=pb, growth=growth)
        calibrated_flow_rows = _build_calibrated_flow_rows(
            years,
            debt=debt,
            primary_balance=pb,
            growth=growth,
            baseline_rows=zero_flow_rows,
        )
        _write_sheet(
            wb,
            sheet_name=short_code,
            country_label=country_label,
            years=years,
            blocks=[
                (ZERO_FLOW_ROW_LAYOUT, zero_flow_rows),
                (CALIBRATED_FLOW_ROW_LAYOUT, calibrated_flow_rows),
            ],
        )

    info = wb.create_sheet(title="INFO")
    info["A1"] = "Data source"
    info["B1"] = "IMF DataMapper API"
    info["A2"] = "Debt indicator"
    info["B2"] = INDICATORS["debt"]
    info["A3"] = "Primary balance indicator"
    info["B3"] = INDICATORS["primary_balance"]
    info["A4"] = "GDP level indicator used"
    info["B4"] = gdp_indicator_used
    info["A5"] = "Note"
    info["B5"] = "Column B is hidden lag year, matching legacy VBA logic. Each country sheet includes zero-flow and calibrated-flow blocks."

    wb.save(inputs.output_path)
    return inputs.output_path
